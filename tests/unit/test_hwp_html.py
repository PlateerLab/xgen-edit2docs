"""HWP 5.0 → HTML (hwp_html) 과 해석 계층(hwp_convert) 보강.

pyhwp(AGPL) 의 hwp5html 을 대신한다. 픽스처는 test_legacy_formats 의 미니 CFB
라이터로 스펙(한글 문서 파일 형식 5.0 revision 1.3, 배포용 문서 revision 1.2)
그대로 조립한다 — 무엇이 스펙의 어느 표를 읽는지가 테스트에 드러난다.

본 제품은 한글과컴퓨터의 글 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다.
"""

from __future__ import annotations

import io
import re
import struct
import zlib

import pytest

from tests.unit.test_legacy_formats import _rec, _tiny_png, _utf16, build_cfb, make_hwp_rich
from xgen_edit2docs.documents.legacy import LegacyConvertError, convert_to_ooxml
from xgen_edit2docs.documents.legacy import hwp_convert as H
from xgen_edit2docs.documents.legacy.hwp_html import ROOT_CLASS, hwp_to_html

# ─────────────────────────────────────────────────────────────
# 픽스처 조립
# ─────────────────────────────────────────────────────────────


def _comp(b: bytes) -> bytes:
    co = zlib.compressobj(6, zlib.DEFLATED, -15)
    return co.compress(b) + co.flush()


def _header(flags: int = 0x1) -> bytes:
    h = bytearray(256)
    h[0:32] = b"HWP Document File".ljust(32, b"\x00")
    struct.pack_into("<I", h, 32, 0x05000000)
    struct.pack_into("<I", h, 36, flags)
    return bytes(h)


def _char_shape(size_pt: float = 10.0, flags: int = 0, color_bgr: int = 0) -> bytes:
    p = bytearray(72)
    struct.pack_into("<i", p, 42, int(size_pt * 100))
    struct.pack_into("<I", p, 46, flags)
    struct.pack_into("<I", p, 52, color_bgr)
    return _rec(0x15, 0, bytes(p))


def _para_shape(flags: int = 1 << 2, left2: int = 0, indent2: int = 0) -> bytes:
    p = bytearray(54)
    struct.pack_into("<I", p, 0, flags)
    struct.pack_into("<i", p, 4, left2)
    struct.pack_into("<i", p, 12, indent2)
    return _rec(0x19, 0, bytes(p))


def _ext(code: int, chid: str) -> bytes:
    """확장 컨트롤 8워드 — ch, ctrl id(4B), 예약 8B, ch."""
    return struct.pack("<H", code) + chid[::-1].encode("ascii") + b"\x00" * 8 + struct.pack("<H", code)


def _para(level: int, text: bytes, spans=((0, 0),), shape: int = 0, extra: bytes = b"") -> bytes:
    hdr = bytearray(16)
    struct.pack_into("<H", hdr, 8, shape)
    cs = b"".join(struct.pack("<II", pos, sid) for pos, sid in spans)
    return (_rec(0x42, level, bytes(hdr)) + _rec(0x43, level + 1, text)
            + _rec(0x44, level + 1, cs) + extra)


def _ctrl(level: int, chid: str, body: bytes = b"") -> bytes:
    return _rec(0x47, level, chid[::-1].encode("ascii") + body)


def _hwp(section: bytes, docinfo: bytes | None = None, *, flags: int = 0x1,
         extra: dict | None = None) -> bytes:
    docinfo = docinfo if docinfo is not None else (_char_shape() + _char_shape(14.0, 0x2)
                                                   + _para_shape())
    streams = {
        "FileHeader": _header(flags),
        "DocInfo": _comp(docinfo) if flags & 0x1 else docinfo,
        "BodyText/Section0": _comp(section) if flags & 0x1 else section,
    }
    streams.update(extra or {})
    return build_cfb(streams)


def _text_of(html: str) -> str:
    body = re.sub(r"<style>.*?</style>", "", html, flags=re.S)
    return re.sub(r"<[^>]+>", "", body)


# ─────────────────────────────────────────────────────────────
# 해석 계층
# ─────────────────────────────────────────────────────────────


class TestStylePositions:
    """PARA_CHAR_SHAPE 위치는 원문 워드 위치(컨트롤 포함)다."""

    def _doc(self) -> bytes:
        # 구역/단 정의(확장 8워드 ×2) 뒤에 "AB" + "CD" — CD 부터 굵게(shape 1).
        text = _ext(2, "secd") + _ext(2, "cold") + _utf16("ABCD")
        return _hwp(_para(0, text, spans=((0, 0), (18, 1))))

    def test_html_puts_the_bold_run_where_the_file_says(self):
        html = hwp_to_html(self._doc())
        assert re.search(r'>AB<span class="c1">CD</span>', html), html

    def test_docx_runs_split_at_the_same_place(self):
        from docx import Document

        d = Document(io.BytesIO(convert_to_ooxml(self._doc(), "hwp")[0]))
        runs = [(r.text, bool(r.bold)) for r in d.paragraphs[0].runs]
        assert runs == [("AB", False), ("CD", True)]


class TestCharacters:
    def test_hyphen_and_the_two_hwp_spaces_are_text(self):
        text = _utf16("4.") + struct.pack("<H", 0x1E) + _utf16("A") + struct.pack(
            "<H", 0x1F) + _utf16("B") + struct.pack("<H", 0x18) + _utf16("C")
        html = hwp_to_html(_hwp(_para(0, text)))
        assert "4. A B-C" in _text_of(html)

    def test_surrogate_pairs_stay_one_character(self):
        pua = "\U000f02b4"  # 한컴 PUA 기호 — UTF-16 서로게이트 쌍
        html = hwp_to_html(_hwp(_para(0, _utf16(f"x{pua}y"))))
        assert f"x{pua}y" in _text_of(html)

    def test_auto_number_goes_into_the_caption(self):
        # 표 캡션 "표 " + 자동 번호(종류 4=표, 번호 3, 뒤 장식 '.')
        atno = struct.pack("<IHHHH", 4, 3, 0, 0, ord("."))
        # 실파일 레벨: PARA(0) > CTRL tbl(1) > [캡션 LIST(2)+문단(2)] TABLE(2) [셀 LIST(2)+문단(2)]
        caption = _rec(0x48, 2, struct.pack("<HHI", 1, 0, 0) + struct.pack("<I", 2)) + _para(
            2, _utf16("표 ") + _ext(0x12, "atno"), extra=_ctrl(3, "atno", atno))
        table = struct.pack("<I", 0) + struct.pack("<HH", 1, 1) + b"\x00" * 12 + b"\x00\x00"
        cell = _rec(0x48, 2, struct.pack("<H", 1) + b"\x00" * 38) + _para(2, _utf16("셀"))
        section = _para(0, _ext(0x0B, "tbl "), extra=_ctrl(1, "tbl ", struct.pack("<I", 1))
                        + caption + _rec(0x4D, 2, table) + cell)
        html = hwp_to_html(_hwp(section))
        text = _text_of(html)
        assert re.search(r"<td[^>]*><p[^>]*>셀</p></td>", html), html
        assert "표 3." in text
        # 캡션 속성 방향 2 = 위 → 표보다 먼저
        assert text.index("표 3.") < text.index("셀")
        from docx import Document

        d = Document(io.BytesIO(convert_to_ooxml(_hwp(section), "hwp")[0]))
        assert any("표 3." in p.text for p in d.paragraphs)

    def test_ruby_main_text_is_kept(self):
        dutmal = struct.pack("<H", 2) + _utf16("본문") + struct.pack("<H", 2) + _utf16("덧말")
        section = _para(0, _utf16("가") + _ext(0x17, "tdut") + _utf16("나"),
                        extra=_ctrl(1, "tdut", dutmal))
        assert "가본문나" in _text_of(hwp_to_html(_hwp(section)))


class TestObjectsInTheLine:
    def _box(self, inline: bool) -> bytes:
        attr = 0x1 if inline else 0x0
        gso = _ctrl(1, "gso ", struct.pack("<I", attr)) + _rec(0x4C, 2, b"\x00" * 36) + _rec(
            0x48, 2, struct.pack("<H", 1) + b"\x00" * 6) + _para(2, _utf16("개화"))
        return _hwp(_para(0, _utf16("지만, ") + _ext(0x0B, "gso ") + _utf16("는 대세"), extra=gso))

    def test_an_inline_text_box_stays_inside_the_sentence(self):
        text = _text_of(hwp_to_html(self._box(inline=True)))
        assert "지만, 개화는 대세" in text

    def test_a_floating_one_comes_after_the_paragraph(self):
        text = _text_of(hwp_to_html(self._box(inline=False)))
        assert text.index("는 대세") < text.index("개화")

    def test_footnote_mark_in_the_line_and_its_text_at_the_section_end(self):
        atno = struct.pack("<IHHHH", 1, 1, 0, 0, ord(")"))
        note = _ctrl(1, "fn  ", b"\x00" * 8) + _rec(0x48, 2, struct.pack("<H", 1) + b"\x00" * 6) \
            + _para(2, _ext(0x12, "atno") + _utf16(" 출처"), extra=_ctrl(3, "atno", atno))
        section = _para(0, _utf16("본문") + _ext(0x11, "fn  ") + _utf16(" 끝"), extra=note) \
            + _para(0, _utf16("다음 문단"))
        html = hwp_to_html(_hwp(section))
        assert '본문<span class="nr">1)</span> 끝' in html
        text = _text_of(html)
        assert text.index("다음 문단") < text.index("1) 출처")

    def test_a_list_hanging_off_a_paragraph_is_not_lost(self):
        """실파일 이형 — 컨트롤 없이 문단 아래 LIST_HEADER + 형제 문단."""
        section = _para(0, _utf16("위"), extra=_rec(0x48, 1, struct.pack("<H", 1) + b"\x00" * 6)
                        + _para(1, _utf16("매달린 내용")))
        assert "매달린 내용" in _text_of(hwp_to_html(_hwp(section)))


class TestStreams:
    def test_a_stored_picture_is_not_inflated(self):
        """BIN_DATA 압축 0x20(압축하지 않음, 표 18) — 한글은 png/jpg 를 날것으로 둔다."""
        bin_data = struct.pack("<HH", 0x21, 1) + struct.pack("<H", 3) + _utf16("png")
        docinfo = (_rec(0x11, 0, struct.pack("<8i", 1, 0, 0, 0, 0, 0, 0, 0))
                   + _rec(0x12, 0, bin_data) + _char_shape() + _para_shape())
        comp = bytearray(36)
        struct.pack_into("<2i", comp, 28, 7200, 7200)
        pic = bytearray(76)
        struct.pack_into("<H", pic, 71, 1)
        gso = _ctrl(1, "gso ", struct.pack("<I", 1)) + _rec(0x4C, 2, bytes(comp)) + _rec(
            0x55, 3, bytes(pic))
        doc = _hwp(_para(0, _ext(0x0B, "gso "), extra=gso), docinfo,
                   extra={"BinData/BIN0001.png": _tiny_png()})  # 문서는 압축, 그림은 날것
        assert "data:image/png;base64," in hwp_to_html(doc)
        from docx import Document

        assert len(Document(io.BytesIO(convert_to_ooxml(doc, "hwp")[0])).inline_shapes) == 1

    def test_a_decompression_bomb_is_refused(self, monkeypatch):
        monkeypatch.setattr(H, "_MAX_STREAM_BYTES", 64)
        monkeypatch.setattr(H._decompress, "__defaults__", (64,))
        with pytest.raises(LegacyConvertError, match="너무 큽니다"):
            hwp_to_html(_hwp(_para(0, _utf16("가" * 200))))

    def test_too_many_records_are_refused(self, monkeypatch):
        """레코드 객체만으로 메모리를 먹는 파일 — 상한에서 멈춘다(실문서 최대 2.6만)."""
        monkeypatch.setattr(H, "_MAX_RECORDS", 5)
        section = b"".join(_para(0, _utf16(f"문단{i}")) for i in range(5))
        with pytest.raises(LegacyConvertError, match="레코드가 너무 많습니다"):
            hwp_to_html(_hwp(section))

    def test_drm_documents_are_refused_by_name(self):
        with pytest.raises(LegacyConvertError, match="DRM"):
            hwp_to_html(_hwp(_para(0, _utf16("x")), flags=0x1 | 0x10))

    def test_landscape_pages_swap_width_and_height(self):
        page = struct.pack("<9I", 59528, 84188, 8504, 8504, 5668, 4252, 0, 0, 0) + struct.pack(
            "<I", 1)
        html = hwp_to_html(_hwp(_rec(0x49, 1, page) + _para(0, _utf16("가로"))))
        assert f"width:{84188 * 96 / 7200:.1f}px" in html


class TestDistributionDocuments:
    """배포용 문서 revision 1.2 §2 — seed → MSVC rand() 배열 → XOR → AES-128 ECB."""

    @staticmethod
    def _distribution_data(key: bytes, options: int, seed: int = 0x12345678) -> bytes:
        rand = H._msvc_rand(seed)
        pattern = bytearray()
        while len(pattern) < 256:
            value = next(rand) & 0xFF
            count = (next(rand) & 0x0F) + 1
            pattern.extend([value] * count)
        merged = bytearray(256)
        struct.pack_into("<I", merged, 0, seed)
        offset = (seed & 0x0F) + 4
        merged[offset:offset + 16] = key
        struct.pack_into("<H", merged, offset + 80, options)
        data = bytes(a ^ b for a, b in zip(merged, pattern[:256], strict=True))
        return struct.pack("<I", seed) + data[4:]  # seed 는 평문 그대로

    def _doc(self, options: int) -> bytes:
        from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes

        key = bytes(range(16))
        body = _comp(_para(0, _utf16("배포용 본문")))
        body += b"\x00" * (-len(body) % 16)
        enc = Cipher(algorithms.AES(key), modes.ECB()).encryptor()
        cipher = enc.update(body) + enc.finalize()
        dist = self._distribution_data(key, options)
        view = _rec(0x1C, 0, dist) + cipher
        return build_cfb({
            "FileHeader": _header(0x1 | 0x4),
            "DocInfo": _comp(_char_shape() + _para_shape()),
            "ViewText/Section0": view,
        })

    def test_the_body_is_decrypted(self):
        assert "배포용 본문" in _text_of(hwp_to_html(self._doc(options=0)))

    def test_copy_and_print_protection_carry_over(self):
        html = hwp_to_html(self._doc(options=0x3))
        assert "user-select:none" in html
        assert f"@media print{{.{ROOT_CLASS}{{display:none;}}}}" in html

    def test_protection_is_known_when_the_file_opens(self):
        with H.HwpFile(self._doc(options=0x1)) as hf:
            assert hf.copy_protected and not hf.print_protected

    def test_copy_protected_text_is_not_handed_out_as_docx(self):
        """DOCX 는 글자를 꺼내 쓰는 길(편집·에이전트 읽기) — 복사 방지 문서는 보기만."""
        with pytest.raises(LegacyConvertError, match="복사가 금지된"):
            convert_to_ooxml(self._doc(options=0x1), "hwp")
        assert "배포용 본문" in _text_of(hwp_to_html(self._doc(options=0x1)))

    def test_an_unprotected_distribution_document_converts(self):
        from docx import Document

        d = Document(io.BytesIO(convert_to_ooxml(self._doc(options=0), "hwp")[0]))
        assert any("배포용 본문" in p.text for p in d.paragraphs)


# ─────────────────────────────────────────────────────────────
# HTML — 충실도와 안전
# ─────────────────────────────────────────────────────────────


@pytest.fixture(scope="module")
def rich_html() -> str:
    return hwp_to_html(make_hwp_rich())


class TestFidelity:
    def test_merges_background_and_alignment(self, rich_html):
        assert 'colspan="2"' in rich_html and 'rowspan="2"' in rich_html
        assert "background-color:#FFCC00" in rich_html
        assert "text-align:center" in rich_html

    def test_run_styles_and_face(self, rich_html):
        assert "line-through" in rich_html and "color:#FF0000" in rich_html
        assert "font-family:'함초롬돋움'" in rich_html
        assert "font-weight:bold" in rich_html

    def test_picture_header_and_borders(self, rich_html):
        assert "data:image/png;base64," in rich_html
        assert "머리말 텍스트" in rich_html
        assert "dashed" in rich_html  # borderfill 2 좌변

    def test_hanging_indent_and_korean_word_breaking(self):
        docinfo = _char_shape() + _para_shape(left2=2000, indent2=-1000)
        html = hwp_to_html(_hwp(_para(0, _utf16("o 내어쓰기 문단")), docinfo))
        assert "margin-left:15pt" in html and "text-indent:-5pt" in html
        assert "word-break:keep-all" in html

    def test_a_header_is_drawn_where_it_is_defined_not_per_section(self):
        head = _ctrl(1, "head", b"\x00" * 8) + _rec(0x48, 2, struct.pack("<H", 1) + b"\x00" * 6) \
            + _para(2, _utf16("머리말 한 번"))
        first = _para(0, _ext(0x10, "head") + _utf16("첫 구역"), extra=head)
        second = _para(0, _utf16("둘째 구역"))
        doc = _hwp(first, extra={"BodyText/Section1": _comp(second)})
        html = hwp_to_html(doc)
        assert "둘째 구역" in html
        assert html.count("머리말 한 번") == 1

    def test_table_border_fill_applies_to_the_table(self):
        bf = bytearray(44)
        for k in range(4):
            bf[2 + k * 6] = 1  # 실선
        docinfo = _char_shape() + _para_shape() + _rec(0x14, 0, bytes(bf))
        # 표 74: 속성4 + 행2 + 열2 + 셀간격2 + 안쪽여백8 + 행 크기 2×1 + BorderFill ID 2
        table = struct.pack("<I", 0) + struct.pack("<HH", 1, 1) + b"\x00" * 10 + struct.pack(
            "<H", 0) + struct.pack("<H", 1)
        cell = _rec(0x48, 2, struct.pack("<H", 1) + b"\x00" * 38) + _para(2, _utf16("셀"))
        section = _para(0, _ext(0x0B, "tbl "), extra=_ctrl(1, "tbl ", struct.pack("<I", 1))
                        + _rec(0x4D, 2, table) + cell)
        html = hwp_to_html(_hwp(section, docinfo))
        assert '<table class="f0"' in html
        assert "border-left:0.1mm solid #000000" in html


class TestSafety:
    def test_text_is_escaped(self):
        html = hwp_to_html(_hwp(_para(0, _utf16("<script>alert(1)</script>&"))))
        assert "<script>" not in html
        assert "&lt;script&gt;alert(1)&lt;/script&gt;&amp;" in html

    def test_a_font_name_cannot_break_out_of_css(self):
        face = b"\x00" + struct.pack("<H", 26) + _utf16("x';}body{display:none}</st")
        docinfo = (_rec(0x11, 0, struct.pack("<8i", 0, 1, 0, 0, 0, 0, 0, 0))
                   + _rec(0x13, 0, face) + _char_shape() + _para_shape())
        html = hwp_to_html(_hwp(_para(0, _utf16("가")), docinfo))
        style = re.search(r"<style>(.*?)</style>", html, re.S).group(1)
        assert "body{display:none}" not in style and "</st" not in style
        assert "font-family:'xbodydisplaynonest'" in style

    def test_every_css_rule_is_scoped_to_the_document(self, rich_html):
        style = re.search(r"<style>(.*?)</style>", rich_html, re.S).group(1)
        selectors = re.findall(r"(?:^|\})\s*([^{}@]+)\{", style)
        assert selectors
        for sel in selectors:
            for part in sel.split(","):
                assert part.strip().startswith(f".{ROOT_CLASS}"), part

    def test_bmp_is_shipped_as_png(self):
        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (4, 4), (255, 0, 0)).save(buf, format="BMP")
        bin_data = struct.pack("<HH", 0x1, 1) + struct.pack("<H", 3) + _utf16("bmp")
        docinfo = (_rec(0x11, 0, struct.pack("<8i", 1, 0, 0, 0, 0, 0, 0, 0))
                   + _rec(0x12, 0, bin_data) + _char_shape() + _para_shape())
        pic = bytearray(76)
        struct.pack_into("<H", pic, 71, 1)
        gso = _ctrl(1, "gso ", struct.pack("<I", 1)) + _rec(0x4C, 2, b"\x00" * 36) + _rec(
            0x55, 3, bytes(pic))
        doc = _hwp(_para(0, _ext(0x0B, "gso "), extra=gso), docinfo,
                   extra={"BinData/BIN0001.bmp": _comp(buf.getvalue())})
        html = hwp_to_html(doc)
        assert "data:image/png;base64," in html and "image/bmp" not in html
